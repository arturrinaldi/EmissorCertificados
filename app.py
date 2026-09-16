import os
import re
import zipfile
import tempfile
import io
import pandas as pd
from docx import Document
from flask import Flask, render_template, request, send_file, jsonify

app = Flask(__name__)
# Get the absolute path to the directory containing this script
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CERTIFICATES_DIR = os.path.join(BASE_DIR, "certificados")

def get_certificate_files():
    files = []
    for f in os.listdir(CERTIFICATES_DIR):
        if f.endswith(".docx") and not f.startswith("~"):
            files.append(f)
    return files

def extract_placeholders(docx_path):
    doc = Document(docx_path)
    text_list = []
    for p in doc.paragraphs:
        text_list.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_list.append(cell.text)
    
    text = '\n'.join(text_list)
    placeholders = set(re.findall(r'«(.*?)»', text))
    return list(placeholders)

def replace_placeholders(docx_path, output_path, data_dict):
    doc = Document(docx_path)
    for p in doc.paragraphs:
        for key, value in data_dict.items():
            token = f"«{key}»"
            if token in p.text:
                for run in p.runs:
                    if token in run.text:
                        run.text = run.text.replace(token, str(value))
                if token in p.text:
                    p.text = p.text.replace(token, str(value))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for key, value in data_dict.items():
                        token = f"«{key}»"
                        if token in p.text:
                            for run in p.runs:
                                if token in run.text:
                                    run.text = run.text.replace(token, str(value))
                            if token in p.text:
                                p.text = p.text.replace(token, str(value))
    doc.save(output_path)

@app.route('/')
def index():
    files = get_certificate_files()
    return render_template('index.html', files=files)

@app.route('/api/placeholders', methods=['GET'])
def get_placeholders():
    filename = request.args.get('file')
    if not filename:
        return jsonify({"error": "File not specified"}), 400
    filepath = os.path.join(CERTIFICATES_DIR, filename)
    if not os.path.exists(filepath):
        return jsonify({"error": "File not found"}), 404
    placeholders = extract_placeholders(filepath)
    return jsonify({"placeholders": placeholders})

@app.route('/download-template', methods=['GET'])
def download_template():
    filename = request.args.get('file')
    if not filename:
        return "File not specified", 400
    filepath = os.path.join(CERTIFICATES_DIR, filename)
    if not os.path.exists(filepath):
        return "File not found", 404
    
    placeholders = extract_placeholders(filepath)
    df = pd.DataFrame(columns=placeholders)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    output.seek(0)
    
    return send_file(
        output,
        as_attachment=True,
        download_name=f"Template_{filename.replace('.docx', '')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/generate', methods=['POST'])
def generate_certificates():
    cert_file = request.form.get('certificate')
    if not cert_file:
        return "Nenhum certificado selecionado.", 400
    
    cert_path = os.path.join(CERTIFICATES_DIR, cert_file)
    if not os.path.exists(cert_path):
        return "Arquivo base não encontrado.", 404

    data_list = []
    
    if 'excel_file' in request.files and request.files['excel_file'].filename:
        excel_file = request.files['excel_file']
        try:
            df = pd.read_excel(excel_file)
            data_list = df.to_dict('records')
        except Exception as e:
            return f"Erro ao ler Excel: {e}", 400
    else:
        import json
        manual_data = request.form.get('manual_data')
        if manual_data:
            try:
                data_list = json.loads(manual_data)
            except Exception as e:
                return f"Erro nos dados manuais: {e}", 400

    if not data_list:
        return "Nenhum dado fornecido para geração.", 400

    memory_file = io.BytesIO()
    with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
        with tempfile.TemporaryDirectory() as tmpdir:
            for i, record in enumerate(data_list):
                out_name = f"Certificado_{i+1}.docx"
                for k, v in record.items():
                    if 'nome' in str(k).lower() and str(v).strip() and str(v).lower() != 'nan':
                        sanitized = "".join([c for c in str(v) if c.isalpha() or c.isdigit() or c==' ']).strip()
                        if sanitized:
                            out_name = f"Certificado_{sanitized}.docx"
                        break
                
                out_path = os.path.join(tmpdir, out_name)
                cleaned_record = {k: ('' if pd.isna(v) else v) for k, v in record.items()}
                replace_placeholders(cert_path, out_path, cleaned_record)
                zf.write(out_path, out_name)
                
    memory_file.seek(0)
    return send_file(
        memory_file,
        as_attachment=True,
        download_name="Certificados_Gerados.zip",
        mimetype="application/zip"
    )

if __name__ == '__main__':
    app.run(debug=True, port=5000)
